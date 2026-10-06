"""Builds the BYOD security remediation workbook for a cloud-only Microsoft 365 tenant.

Outputs (next to this script):
  BYOD-Security-Remediation.numbers   Apple Numbers
  BYOD-Security-Remediation.xlsx      Excel (also opens in Numbers)

Requires: openpyxl, numbers-parser   (pip install openpyxl numbers-parser)
"""

from pathlib import Path

HERE = Path(__file__).parent
NAME = "BYOD-Security-Remediation"

COLUMNS = [
    "Priority",
    "Criticality",
    "Area",
    "Issue",
    "BYOD impact and risk",
    "Remediation steps",
    "Admin console",
    "Licence needed",
    "How to verify",
    "Effort",
    "Owner",
    "Status",
]

# Ordered by criticality. Licence notes: Business Premium includes Entra ID P1, Intune and
# Defender for Business; Entra ID P2 / Defender for Cloud Apps come with E5 or as add-ons.
ITEMS = [
    {
        "Criticality": "Critical",
        "Area": "Foundation",
        "Issue": "Licensing does not include the controls BYOD depends on",
        "BYOD impact and risk": (
            "Every meaningful BYOD control in Microsoft 365 needs Entra ID P1 (Conditional Access) and "
            "Intune (app protection / device compliance). Tenants on Business Basic/Standard or E1/E3 "
            "without EMS can only use Security Defaults, which enforces MFA but cannot tell a personal "
            "device from a managed one, restrict downloads or wipe corporate data. Without these licences "
            "personal devices get exactly the same access as company-managed ones, and nothing else in "
            "this list can be fully implemented."
        ),
        "Remediation steps": (
            "1. Microsoft 365 admin center > Billing > Licenses: note which products you own and how many "
            "are assigned.\n"
            "2. Microsoft 365 admin center > Users > Active users > Filter > Unlicensed users: identify "
            "anyone without a licence.\n"
            "3. Microsoft 365 admin center > Billing > Purchase services: buy Microsoft 365 Business Premium "
            "(up to 300 users) or E3 + EMS E3 / E5. Add Entra ID P2 or Defender for Cloud Apps only if you "
            "want items 5 (risk policies) and 9 (browser session control).\n"
            "4. Microsoft 365 admin center > Users > Active users > select users > Manage product licenses > "
            "Replace: assign the new licence to everyone who will use a personal device. Unlicensed users "
            "are not covered by Intune app protection."
        ),
        "Admin console": "Microsoft 365 admin center (admin.microsoft.com) > Billing; Users > Active users",
        "Licence needed": "Business Premium, or E3 + EMS E3 / E5",
        "How to verify": "Microsoft 365 admin center > Billing > Licenses shows Business Premium (or E3+EMS/E5) assigned to every active user; the Unlicensed users filter returns nobody who uses company data.",
        "Effort": "Low",
    },
    {
        "Criticality": "Critical",
        "Area": "Identity",
        "Issue": "MFA not enforced for every user and legacy authentication still allowed",
        "BYOD impact and risk": (
            "Personal devices are outside IT's control, so the user's password is the only thing standing "
            "between an attacker and the mailbox. Passwords reused on personal sites or captured by malware "
            "on a home PC give full access to Exchange, SharePoint and Teams from anywhere. Legacy protocols "
            "(POP, IMAP, SMTP AUTH with Basic auth, older ActiveSync clients) cannot do MFA and are the main "
            "target of password-spray attacks."
        ),
        "Remediation steps": (
            "1. No Conditional Access licence: Entra admin center > Entra ID > Overview > Properties > Manage "
            "security defaults > Enabled. (Turn it off again once you build the policies below.)\n"
            "2. With Entra ID P1: Entra admin center > Conditional Access > Policies > New policy from "
            "template > 'Require multifactor authentication for all users' and 'Block legacy "
            "authentication'. Set each to Report-only.\n"
            "3. In each policy: Users > Exclude > add two cloud-only break-glass admin accounts.\n"
            "4. After 1-2 weeks, Entra admin center > Sign-in logs > filter Conditional Access = Report-only: "
            "failure to check who would be blocked, fix them, then set the policies to On.\n"
            "5. Entra admin center > Authentication methods > Policies: enable Passkey (FIDO2) and Microsoft "
            "Authenticator; then Conditional Access > Authentication strengths > use 'Phishing-resistant MFA' "
            "in a policy targeting admin roles.\n"
            "6. Exchange admin center > Settings > Mail flow > tick 'Turn off SMTP AUTH protocol for your "
            "organization'.\n"
            "7. For users who do not need POP/IMAP: Microsoft 365 admin center > Users > Active users > "
            "select user > Mail > Manage email apps > untick POP, IMAP and Authenticated SMTP > Save changes."
        ),
        "Admin console": "Entra admin center (entra.microsoft.com) > Conditional Access, Authentication methods; Exchange admin center > Settings; Microsoft 365 admin center > Active users",
        "Licence needed": "Free (Security Defaults) / Entra ID P1 (Conditional Access)",
        "How to verify": "Entra admin center > Sign-in logs > filter Client app = the legacy client types: only failures appear. Entra admin center > Authentication methods > User registration details: every user is MFA capable.",
        "Effort": "Medium",
    },
    {
        "Criticality": "Critical",
        "Area": "Access control",
        "Issue": "No Conditional Access rule distinguishes personal from managed devices",
        "BYOD impact and risk": (
            "Without device-aware Conditional Access, a personal laptop or phone receives full desktop-client "
            "access: Outlook can cache the whole mailbox, OneDrive can sync every file, and Teams can download "
            "attachments to unencrypted local storage. When the device is lost, sold, shared with family or "
            "infected, that data leaves the organisation with no way to recall it."
        ),
        "Remediation steps": (
            "Decide the model first: phones/tablets = app protection without enrolment; personal Windows/Mac "
            "= browser only (or Edge with Windows MAM); full desktop apps only on Intune-compliant devices.\n"
            "1. Entra admin center > Conditional Access > Policies > New policy 'BYOD - mobile apps need app "
            "protection': Users = All users (exclude break-glass); Target resources = Office 365; Conditions > "
            "Device platforms = iOS, Android; Grant = 'Require app protection policy'. (The older 'Require "
            "approved client app' control is being retired.)\n"
            "2. New policy 'Desktop apps need compliant device': Device platforms = Windows, macOS; Conditions "
            "> Client apps = Mobile apps and desktop clients; Grant = 'Require device to be marked as "
            "compliant'.\n"
            "3. New policy 'Browser on unmanaged devices': Client apps = Browser; Session = 'Use app enforced "
            "restrictions' (works with the SharePoint/Outlook settings in item 4).\n"
            "4. New policy 'Block unknown platforms': Device platforms > Include Any device > Exclude Android, "
            "iOS, Windows, macOS; Grant = Block access.\n"
            "5. Create every policy as Report-only, test with Conditional Access > Policies > What If, then "
            "switch to On."
        ),
        "Admin console": "Entra admin center (entra.microsoft.com) > Conditional Access > Policies",
        "Licence needed": "Entra ID P1 + Intune",
        "How to verify": "Entra admin center > Conditional Access > What If: a user on iOS gets 'Require app protection policy', on an unmanaged Windows desktop client gets 'Require compliant device'. Sign-in logs > Conditional Access tab shows the expected policy applied.",
        "Effort": "High",
    },
    {
        "Criticality": "Critical",
        "Area": "Data protection",
        "Issue": "Corporate data can be saved, synced or copied into personal apps and storage",
        "BYOD impact and risk": (
            "On a personal phone, a user can open an email attachment and save it to personal iCloud or "
            "Google Drive, paste customer data into a personal messaging app, or screenshot it. On a home PC, "
            "files download to an unencrypted disk that IT cannot wipe. This is the most common way data "
            "leaves organisations that allow BYOD, often unintentionally, and it creates breach-notification "
            "exposure under GDPR and similar laws."
        ),
        "Remediation steps": (
            "1. Intune admin center > Apps > App protection policies > Create policy > iOS/iPadOS (repeat for "
            "Android). Apps: Target to 'All Microsoft apps' (or select Outlook, Teams, OneDrive, SharePoint, "
            "Word, Excel, PowerPoint, Edge). Assignments: All users.\n"
            "2. Data protection page: Backup org data to iTunes/iCloud/Google = Block; Send org data to other "
            "apps = Policy managed apps; Save copies of org data = Block, Allow user to save copies to = "
            "OneDrive for Business, SharePoint; Restrict cut, copy and paste = Policy managed apps with paste "
            "in; Encrypt org data = Require; Screen capture (Android) = Block.\n"
            "3. Access requirements page: PIN for access = Require (biometrics allowed); Recheck after "
            "(minutes of inactivity) = 30.\n"
            "4. SharePoint admin center > Policies > Access control > Unmanaged devices > 'Allow limited, "
            "web-only access' > Save. (This also creates matching Conditional Access policies in Entra; "
            "review them there.)\n"
            "5. Outlook on the web: Exchange has no admin-center switch for read-only attachments. Use "
            "Defender portal > Cloud apps > Policies > Policy management > Create policy > Session policy > "
            "'Block download based on real-time content inspection' scoped to unmanaged devices (needs "
            "Defender for Cloud Apps). Without that licence, the only option is one Exchange PowerShell "
            "command (Set-OwaMailboxPolicy ... -ConditionalAccessPolicy ReadOnly).\n"
            "6. Personal Windows PCs: Intune admin center > Apps > App protection policies > Create policy > "
            "Windows (Microsoft Edge), then in Entra a CA policy for Windows + Browser with Grant 'Require "
            "app protection policy'.\n"
            "7. The OneDrive 'sync only on domain-joined PCs' setting needs on-premises AD and does not apply "
            "to a cloud-only tenant; the compliant-device policy in item 3 controls sync instead."
        ),
        "Admin console": "Intune admin center (intune.microsoft.com) > Apps > App protection policies; SharePoint admin center > Policies > Access control; Defender portal > Cloud apps",
        "Licence needed": "Intune + Entra ID P1 (Defender for Cloud Apps for step 5)",
        "How to verify": "On a test personal phone, 'Save to Files' from Outlook and pasting mail text into a personal app are both blocked. In a browser on an unmanaged PC, SharePoint shows the 'download disabled' banner. Intune admin center > Apps > Monitor > App protection status lists the test user.",
        "Effort": "High",
    },
    {
        "Criticality": "Critical",
        "Area": "Identity",
        "Issue": "Session and token theft from personal devices (AiTM phishing, info-stealer malware)",
        "BYOD impact and risk": (
            "Home computers frequently lack EDR and run unvetted software. Info-stealer malware lifts browser "
            "cookies and refresh tokens, and adversary-in-the-middle phishing kits proxy the MFA prompt. Either "
            "way the attacker replays the session from their own machine, bypassing MFA entirely, and "
            "typically sets up inbox forwarding rules for business email compromise."
        ),
        "Remediation steps": (
            "1. Entra admin center > Authentication methods > Policies: enable Passkey (FIDO2) and roll out "
            "passkeys in Microsoft Authenticator; phishing-resistant methods defeat AiTM proxies.\n"
            "2. Entra admin center > Conditional Access > New policy 'Unmanaged device sessions': Conditions > "
            "Filter for devices > exclude compliant devices (device.isCompliant -eq True); Session > Sign-in "
            "frequency = 8 hours and Persistent browser session = Never persistent.\n"
            "3. In the same policy, Session > Customize continuous access evaluation: leave enabled; if you "
            "use named locations choose 'Strictly enforce location policies'.\n"
            "4. Where your platforms support it, Session > 'Require token protection for sign-in sessions' for "
            "Exchange Online, SharePoint and Teams (check Microsoft's current platform list first).\n"
            "5. With Entra ID P2: new CA policies using Conditions > Sign-in risk (High, Medium) = Require MFA, "
            "and Conditions > User risk (High) = Require password change.\n"
            "6. Compromise playbook in the consoles: Microsoft 365 admin center > Active users > user > "
            "Account > 'Sign out of all sessions' (or Entra admin center > Users > user > Revoke sessions); "
            "Reset password; Exchange admin center > Mailboxes > user > Mail flow settings > check Email "
            "forwarding; Purview portal > Audit > search activities 'New-InboxRule' / 'Set-InboxRule' for that "
            "user; Exchange admin center > Reports > Mail flow > Auto forwarded messages."
        ),
        "Admin console": "Entra admin center > Conditional Access, Authentication methods, Users; Microsoft 365 admin center; Exchange admin center; Purview portal > Audit",
        "Licence needed": "Entra ID P1 (P2 for risk policies)",
        "How to verify": "Entra admin center > Sign-in logs > a sign-in from an unmanaged device shows sign-in frequency applied. Run the playbook on a test account: the user is signed out of Outlook/Teams within minutes.",
        "Effort": "Medium",
    },
    {
        "Criticality": "High",
        "Area": "Data protection",
        "Issue": "No way to remove corporate data from lost, stolen or leavers' personal devices",
        "BYOD impact and risk": (
            "When an employee leaves or loses a phone, cached mail, Teams chats and downloaded files remain on "
            "a device the organisation does not own. A full device wipe is usually unacceptable (or unlawful) "
            "on personal property, so without app-level wipe the data simply stays there indefinitely."
        ),
        "Remediation steps": (
            "1. Rely on app protection (item 4) so corporate data lives only inside managed apps.\n"
            "2. Intune admin center > Apps > App protection policies > open each policy > Properties > "
            "Conditional launch > Edit: Offline grace period = Wipe data after 90 days.\n"
            "3. Lost device or leaver: Intune admin center > Apps > App selective wipe > Create wipe request > "
            "choose the user and device.\n"
            "4. Microsoft 365 admin center > Active users > user > Block sign-in, then Account > 'Sign out of "
            "all sessions'.\n"
            "5. If native mail apps are still allowed: Exchange admin center > Recipients > Mailboxes > select "
            "user > Others > Manage mobile devices > select the device > 'Account only remote wipe device' "
            "(removes the mailbox only, leaves personal data).\n"
            "6. Add these steps to the HR leaver checklist and tell users to report a lost device within hours."
        ),
        "Admin console": "Intune admin center > Apps > App selective wipe; Microsoft 365 admin center > Active users; Exchange admin center > Mailboxes",
        "Licence needed": "Intune",
        "How to verify": "Create a selective wipe for a test phone: Intune admin center > Apps > App selective wipe shows status Done; company accounts and files disappear from Outlook/OneDrive/Teams while personal photos and apps remain.",
        "Effort": "Low",
    },
    {
        "Criticality": "High",
        "Area": "Device health",
        "Issue": "Out-of-date, jailbroken/rooted or malware-infected personal devices",
        "BYOD impact and risk": (
            "Personal devices often run old OS versions with known exploits, may be jailbroken or rooted "
            "(which defeats app sandboxing and app protection), or share a PC with family members who install "
            "anything. Such devices can leak data even inside managed apps."
        ),
        "Remediation steps": (
            "1. Intune admin center > Apps > App protection policies > open each policy > Properties > "
            "Conditional launch > Edit. Device conditions: Jailbroken/rooted devices = Block access; Min OS "
            "version = current major minus one, action Block access; Play integrity verdict (Android) = Basic "
            "integrity & device certification, Block access. App conditions: Min app version where needed.\n"
            "2. Optional mobile threat defence: Intune admin center > Endpoint security > Microsoft Defender "
            "for Endpoint > turn on 'Connect Android/iOS devices to Microsoft Defender for Endpoint' for app "
            "protection policy evaluation; then in Conditional launch add Max allowed device threat level = "
            "Low.\n"
            "3. For personal devices you enrol (iOS User Enrollment, Android Work Profile): Intune admin "
            "center > Devices > Compliance > Create policy: minimum OS, require passcode and encryption.\n"
            "4. Intune admin center > Devices > Compliance > Compliance settings: 'Mark devices with no "
            "compliance policy assigned as' = Not compliant."
        ),
        "Admin console": "Intune admin center (intune.microsoft.com) > Apps > App protection policies; Devices > Compliance; Endpoint security",
        "Licence needed": "Intune (Defender for Endpoint optional)",
        "How to verify": "Intune admin center > Apps > Monitor > App protection status > App protection report: no users on blocked OS versions. A test device below the minimum OS is blocked on launch.",
        "Effort": "Medium",
    },
    {
        "Criticality": "High",
        "Area": "Email",
        "Issue": "Native mail apps (iOS Mail, Samsung/Gmail) and ActiveSync bypass app protection",
        "BYOD impact and risk": (
            "Built-in mail apps use Exchange ActiveSync and cannot be governed by Intune app protection: "
            "attachments can be saved anywhere and mail is stored outside any managed container. Any "
            "ActiveSync client can connect by default."
        ),
        "Remediation steps": (
            "1. See who uses what: Microsoft 365 admin center > Reports > Usage > Exchange > Email apps usage.\n"
            "2. Announce the move to Outlook for iOS/Android, with a date.\n"
            "3. Entra admin center > Conditional Access > New policy 'Block ActiveSync': Conditions > Client "
            "apps > Exchange ActiveSync clients; Target resources = Office 365 Exchange Online; Grant = Block "
            "access. Report-only first, then On.\n"
            "4. Defence in depth: Exchange admin center > Mobile > Device access > Edit: when a device that "
            "is not managed by a rule connects = Quarantine; then add a device access rule allowing Outlook.\n"
            "5. Users who need no mobile mail at all: Microsoft 365 admin center > Active users > user > Mail > "
            "Manage email apps > untick Exchange ActiveSync."
        ),
        "Admin console": "Entra admin center > Conditional Access; Exchange admin center > Mobile > Device access; Microsoft 365 admin center > Reports",
        "Licence needed": "Entra ID P1 (Exchange controls are included with Exchange Online)",
        "How to verify": "Adding the work account to iOS Mail fails. Microsoft 365 admin center > Reports > Usage > Email apps usage shows Outlook for mobile only after cleanup.",
        "Effort": "Low",
    },
    {
        "Criticality": "High",
        "Area": "Data protection",
        "Issue": "Unmanaged browsers and desktop apps on personal Windows and Mac computers",
        "BYOD impact and risk": (
            "A personal PC with any browser can open Outlook on the web, SharePoint and Teams. Without session "
            "controls the user can download, print and sync freely, and malware on that PC can read anything "
            "opened. Office desktop apps and the OneDrive sync client on such PCs create full local copies."
        ),
        "Remediation steps": (
            "1. Desktop apps blocked on non-compliant Windows/macOS: the 'Desktop apps need compliant device' "
            "policy from item 3.\n"
            "2. Browser access limited: the SharePoint 'Allow limited, web-only access' setting and the "
            "browser policy from items 3-4.\n"
            "3. Preferred on personal Windows: Intune admin center > Apps > App protection policies > Create "
            "policy > Windows (Microsoft Edge) with Data protection > Receive data from / Send data to = "
            "Org data sources only, Print org data = Block. Then Entra admin center > Conditional Access > New "
            "policy: Device platforms = Windows, Client apps = Browser, Grant = Require app protection "
            "policy. Users get the Edge work profile.\n"
            "4. With Defender for Cloud Apps: Defender portal > Cloud apps > Policies > Policy management > "
            "Create policy > Session policy: Session control type = Control file download (with inspection), "
            "Activity source = device tag not Intune compliant, Action = Block (also add policies for print "
            "and copy/paste as needed).\n"
            "5. Explain to users that personal PCs are browser-only by design."
        ),
        "Admin console": "Entra admin center > Conditional Access; Intune admin center > App protection policies (Windows); Defender portal (security.microsoft.com) > Cloud apps",
        "Licence needed": "Entra ID P1 + Intune; Defender for Cloud Apps for session policies",
        "How to verify": "On a personal PC: Outlook desktop sign-in is blocked, SharePoint shows the 'download disabled' banner, and (with Cloud Apps) a download attempt shows the block page. Defender portal > Cloud apps > Activity log records it.",
        "Effort": "Medium",
    },
    {
        "Criticality": "High",
        "Area": "Identity",
        "Issue": "Uncontrolled device registration and enrolment",
        "BYOD impact and risk": (
            "By default any user can register (Entra registered) or join devices without extra verification, "
            "and enrol personal Windows PCs into Intune. An attacker with a stolen password can register their "
            "own device, which may then satisfy weaker device-based policies or persist access. Stale "
            "registrations clutter inventory and hide real devices."
        ),
        "Remediation steps": (
            "1. Entra admin center > Devices > Device settings: Maximum number of devices per user = 10 (or "
            "fewer).\n"
            "2. Entra admin center > Conditional Access > New policy: Target resources > User actions > "
            "'Register or join devices'; Grant = Require multifactor authentication. Then in Devices > Device "
            "settings set 'Require Multifactor Authentication to register or join devices' = No (the CA "
            "policy replaces it).\n"
            "3. Intune admin center > Devices > Enrollment > Device platform restriction > edit the default "
            "Windows and macOS restrictions: Personally owned devices = Block (if personal PCs are browser "
            "only). For iOS allow User Enrollment; for Android allow Work Profile only.\n"
            "4. Monthly: Entra admin center > Devices > All devices > Columns: Activity > filter devices "
            "inactive 90+ days > Disable; delete after a further grace period."
        ),
        "Admin console": "Entra admin center > Devices > Device settings, All devices; Conditional Access; Intune admin center > Devices > Enrollment",
        "Licence needed": "Entra ID P1 + Intune",
        "How to verify": "Registering a new device prompts for MFA. Enrolling a personal Windows PC into Intune is refused. Entra admin center > Devices > All devices shows no active devices older than 90 days.",
        "Effort": "Low",
    },
    {
        "Criticality": "High",
        "Area": "Identity",
        "Issue": "Users can grant third-party apps access to company data (OAuth consent)",
        "BYOD impact and risk": (
            "On personal devices users install productivity apps, mail clients and AI tools that ask to "
            "'Sign in with Microsoft' and request Mail.Read or Files.Read.All. Once consented, the app reads "
            "data directly through Microsoft Graph from its own servers, bypassing every device control. "
            "Consent phishing uses the same path."
        ),
        "Remediation steps": (
            "1. Entra admin center > Enterprise apps > Consent and permissions > User consent settings: "
            "'Allow user consent for apps from verified publishers, for selected permissions' (or 'Do not "
            "allow user consent').\n"
            "2. Same page > Admin consent settings: Users can request admin consent = Yes; choose reviewers.\n"
            "3. Entra admin center > Enterprise apps > All applications > Application type = All applications: "
            "open each unfamiliar app > Permissions; remove apps you do not recognise (Properties > Delete).\n"
            "4. With Defender for Cloud Apps: Defender portal > Cloud apps > App governance: turn on and review "
            "the predefined OAuth app policies."
        ),
        "Admin console": "Entra admin center > Enterprise apps > Consent and permissions; Defender portal > Cloud apps > App governance",
        "Licence needed": "Free (Defender for Cloud Apps optional)",
        "How to verify": "A test user who tries to consent to an unverified app requesting Mail.Read sees 'Approval required'. Entra admin center > Enterprise apps > Admin consent requests receives the request.",
        "Effort": "Low",
    },
    {
        "Criticality": "Medium",
        "Area": "Monitoring",
        "Issue": "Little visibility of which personal devices access what",
        "BYOD impact and risk": (
            "Without device-aware logging you cannot answer 'which unmanaged devices accessed customer data "
            "last week?' during an incident. Entra ID Free keeps sign-in logs for only 7 days, P1/P2 for 30 "
            "days, which is often shorter than the time it takes to discover a compromise."
        ),
        "Remediation steps": (
            "1. Microsoft Purview portal > Solutions > Audit: if a banner says 'Start recording user and admin "
            "activity', select it (auditing is on by default in most tenants).\n"
            "2. Entra admin center > Monitoring & health > Diagnostic settings > Add diagnostic setting: send "
            "SignInLogs, NonInteractiveUserSignInLogs and AuditLogs to a Log Analytics workspace (Azure "
            "subscription) with 90+ days retention.\n"
            "3. Entra admin center > Monitoring & health > Workbooks > 'Conditional Access insights and "
            "reporting': review unmanaged-device sign-ins monthly.\n"
            "4. Intune admin center > Apps > Monitor > App protection status: review monthly.\n"
            "5. Defender portal > Email & collaboration > Policies & rules > Alert policy: make sure alerts for "
            "'Suspicious email forwarding activity' and 'Creation of forwarding/redirect rule' are on and "
            "notify an admin mailbox."
        ),
        "Admin console": "Purview portal (purview.microsoft.com) > Audit; Entra admin center > Monitoring & health; Defender portal > Policies & rules",
        "Licence needed": "Entra ID P1; Azure subscription for Log Analytics",
        "How to verify": "Purview > Audit search returns events from today. A sign-in from a test unmanaged device can still be found in Log Analytics 60 days later.",
        "Effort": "Medium",
    },
    {
        "Criticality": "Medium",
        "Area": "Data protection",
        "Issue": "Sensitive data is not classified, so protection cannot follow it onto personal devices",
        "BYOD impact and risk": (
            "Device controls protect the container, but once data leaves (forwarded email, a file shared "
            "externally, a screenshot on another device) only content-level protection still applies. "
            "Without sensitivity labels and DLP, a payroll spreadsheet and a lunch menu are treated the same."
        ),
        "Remediation steps": (
            "1. Microsoft Purview portal > Solutions > Information protection > Sensitivity labels > Create a "
            "label: Public, General, Confidential, Highly Confidential.\n"
            "2. For Confidential and Highly Confidential: Items > Control access (encryption) > assign "
            "permissions to your organisation; for Highly Confidential set 'Allow offline access' = Never.\n"
            "3. Information protection > Label publishing policies > Publish labels to all users; optionally "
            "set a default label.\n"
            "4. Purview portal > Solutions > Data loss prevention > Policies > Create policy: use a template "
            "(e.g. Financial / Privacy for your country), locations Exchange, SharePoint, OneDrive, Teams, "
            "action Block sharing outside the organisation with policy tips. Start in simulation mode."
        ),
        "Admin console": "Microsoft Purview portal (purview.microsoft.com) > Information protection; Data loss prevention",
        "Licence needed": "Business Premium / E3 (manual labels, DLP); E5 for auto-labelling",
        "How to verify": "An unauthorised user cannot open a Highly Confidential test file. Emailing a test credit-card number externally triggers the DLP policy tip and appears under Purview > Data loss prevention > Alerts.",
        "Effort": "High",
    },
    {
        "Criticality": "Medium",
        "Area": "Governance",
        "Issue": "No written BYOD policy covering acceptable use, privacy and what IT can see or wipe",
        "BYOD impact and risk": (
            "Without a signed policy, enforcing controls on personal property invites disputes and possible "
            "legal exposure (employee privacy, unlawful wipe of personal data). Users who do not understand "
            "the rules work around them, for example by emailing files to personal accounts. eDiscovery and "
            "legal holds may also need to reach data that sits on personal devices."
        ),
        "Remediation steps": (
            "1. Write the policy with HR and legal: eligible devices, required apps, minimum OS, PIN/biometric, "
            "report loss within 24 hours, consent to selective wipe, what IT can and cannot see.\n"
            "2. Entra admin center > Conditional Access > Terms of use > New terms: upload the policy PDF, "
            "Require users to expand = On.\n"
            "3. Entra admin center > Conditional Access > New policy 'Accept BYOD terms': All users, Office 365, "
            "Grant = the terms of use you created.\n"
            "4. Intune admin center > Tenant administration > Customization > edit: add privacy statement URL, "
            "support contact and the Company Portal message explaining what IT can see.\n"
            "5. Review the policy every year and publish a new terms-of-use version when it changes."
        ),
        "Admin console": "Entra admin center > Conditional Access > Terms of use; Intune admin center > Tenant administration > Customization",
        "Licence needed": "Entra ID P1 (terms of use)",
        "How to verify": "Entra admin center > Conditional Access > Terms of use > select the terms > View accepted/declined shows every BYOD user accepted the current version.",
        "Effort": "Medium",
    },
    {
        "Criticality": "Medium",
        "Area": "Data protection",
        "Issue": "Personal and work accounts mixed in the same apps and browser",
        "BYOD impact and risk": (
            "On a personal device the same Outlook, OneDrive or Edge app often holds both a personal account "
            "and the work account. Users can move content between them, or sign in to other organisations' "
            "tenants and upload company data there."
        ),
        "Remediation steps": (
            "1. App protection (item 4) already separates accounts inside Outlook/OneDrive/Office: confirm "
            "'Save copies of org data' = Block with only OneDrive for Business and SharePoint allowed.\n"
            "2. Personal Windows: the Edge work profile created by Windows MAM (item 9) keeps work browsing "
            "separate.\n"
            "3. Entra admin center > External Identities > Cross-tenant access settings > Default settings > "
            "Tenant restrictions: block access to other tenants for your users. Note it is only enforced on "
            "managed devices or through Global Secure Access, not on unmanaged personal devices.\n"
            "4. SharePoint admin center > Policies > Sharing: set external sharing to 'Existing guests' or "
            "'New and existing guests' and limit by domain (More external sharing settings)."
        ),
        "Admin console": "Intune admin center > App protection policies; Entra admin center > External Identities; SharePoint admin center > Policies > Sharing",
        "Licence needed": "Intune + Entra ID P1",
        "How to verify": "On a test phone, saving a work attachment to the personal OneDrive account in the same app is blocked.",
        "Effort": "Medium",
    },
    {
        "Criticality": "Medium",
        "Area": "People",
        "Issue": "Users not trained on BYOD risks and how to report incidents",
        "BYOD impact and risk": (
            "Technical controls fail if users install apps from unofficial stores, ignore OS updates, approve "
            "unexpected MFA prompts or do not report a lost phone. Personal devices are also used for personal "
            "email and social media, which raises phishing exposure."
        ),
        "Remediation steps": (
            "1. Run short onboarding training: installing Outlook/Teams from the official stores, what app "
            "protection does, how to report loss or compromise.\n"
            "2. Defender portal > Email & collaboration > Attack simulation training > Simulations > Launch a "
            "simulation (Defender for Office 365 Plan 2).\n"
            "3. Defender portal > Settings > Email & collaboration > User reported settings: enable the "
            "Microsoft Report button in Outlook and set a reporting mailbox.\n"
            "4. Publish a one-page 'lost device' procedure with the helpdesk contact; link it in the Company "
            "Portal (Intune admin center > Tenant administration > Customization)."
        ),
        "Admin console": "Defender portal (security.microsoft.com) > Attack simulation training; Settings > Email & collaboration",
        "Licence needed": "Defender for Office 365 Plan 2 for simulations (optional)",
        "How to verify": "Defender portal > Attack simulation training > Reports: completion and click rates tracked quarterly. Reported messages appear under Actions & submissions > Submissions.",
        "Effort": "Low",
    },
    {
        "Criticality": "Low",
        "Area": "Network",
        "Issue": "Access from untrusted networks (public Wi-Fi, home routers)",
        "BYOD impact and risk": (
            "Microsoft 365 traffic is always TLS-encrypted, so network eavesdropping is a low risk. The "
            "remaining exposure is captive-portal phishing and sign-ins from unexpected countries, which "
            "location-aware policies can reduce."
        ),
        "Remediation steps": (
            "1. Entra admin center > Conditional Access > Named locations > + Countries location: add the "
            "countries you operate in.\n"
            "2. Conditional Access > New policy: Conditions > Locations > Include Any location, Exclude your "
            "countries location; Grant = Block access (or require phishing-resistant MFA). Exclude break-glass "
            "accounts and remember travelling staff.\n"
            "3. Keep MFA and session controls (items 2 and 5) as the main protection.\n"
            "4. Optional: Entra admin center > Global Secure Access to add compliant-network checks."
        ),
        "Admin console": "Entra admin center > Conditional Access > Named locations, Policies",
        "Licence needed": "Entra ID P1",
        "How to verify": "Conditional Access > What If with a blocked country selected returns Block access; sign-in logs show the policy blocking a test sign-in through a foreign VPN.",
        "Effort": "Low",
    },
    {
        "Criticality": "Low",
        "Area": "Resilience",
        "Issue": "Accidental deletion or ransomware on a synced personal device",
        "BYOD impact and risk": (
            "If a personal PC with OneDrive sync is hit by ransomware or a user deletes a folder, the change "
            "syncs to the cloud. Microsoft 365 has recovery tools, but retention is limited unless configured."
        ),
        "Remediation steps": (
            "1. Prefer blocking sync on unmanaged PCs (items 3 and 4) so this cannot occur.\n"
            "2. Know the recovery paths: user's OneDrive > Settings > Restore your OneDrive (30 days); "
            "SharePoint site > Recycle bin (93 days).\n"
            "3. Microsoft Purview portal > Solutions > Data lifecycle management > Policies > Retention "
            "policies > New retention policy: Exchange mailboxes, SharePoint sites, OneDrive accounts; retain "
            "for the period your business needs.\n"
            "4. Microsoft 365 admin center > Settings > Microsoft 365 Backup: set up billing and backup "
            "policies for critical OneDrive accounts, SharePoint sites and mailboxes."
        ),
        "Admin console": "Purview portal > Data lifecycle management; Microsoft 365 admin center > Settings > Microsoft 365 Backup",
        "Licence needed": "Included (retention); Microsoft 365 Backup is pay-as-you-go",
        "How to verify": "Restore a deleted test folder from a pilot user's OneDrive; Purview > Data lifecycle management > Retention policies shows the policy On for all three locations.",
        "Effort": "Low",
    },
]

CONSOLES = [
    ("Microsoft 365 admin center", "https://admin.microsoft.com", "Licences, users, email apps per user, sign out of sessions, usage reports, Microsoft 365 Backup", "Global Reader to view; User Administrator / Billing Administrator to change"),
    ("Microsoft Entra admin center", "https://entra.microsoft.com", "Conditional Access, authentication methods, device settings, enterprise app consent, terms of use, sign-in logs", "Conditional Access Administrator, Authentication Policy Administrator, Cloud Application Administrator"),
    ("Microsoft Intune admin center", "https://intune.microsoft.com", "App protection policies, selective wipe, compliance, enrolment restrictions, Company Portal customisation", "Intune Administrator"),
    ("Exchange admin center", "https://admin.exchange.microsoft.com", "SMTP AUTH, mobile device access, mailbox forwarding, account-only wipe, mail flow reports", "Exchange Administrator"),
    ("SharePoint admin center", "https://<tenant>-admin.sharepoint.com (or Microsoft 365 admin center > Admin centers > SharePoint)", "Unmanaged device access, external sharing", "SharePoint Administrator"),
    ("Microsoft Defender portal", "https://security.microsoft.com", "Cloud Apps session policies and app governance, alert policies, attack simulation, user-reported settings", "Security Administrator"),
    ("Microsoft Purview portal", "https://purview.microsoft.com", "Audit, sensitivity labels, DLP, retention", "Compliance Administrator"),
]

CRIT_FILL = {  # background, text
    "Critical": ("C00000", "FFFFFF"),
    "High": ("F4B183", "000000"),
    "Medium": ("FFE699", "000000"),
    "Low": ("C6E0B4", "000000"),
}

WIDTHS = [8, 11, 14, 34, 60, 95, 36, 24, 44, 9, 14, 13]

README_ROWS = [
    ("BYOD security remediation plan - Microsoft 365 cloud-only tenant", ""),
    ("", ""),
    ("Purpose", "Issues that bring-your-own-device creates for a cloud-only Microsoft 365 tenant, ordered by criticality, with remediation steps."),
    ("How to use", "Work top to bottom on the 'Remediation Plan' sheet. Items 1-5 (Critical) should be in place before allowing personal devices access to company data."),
    ("Admin consoles", "All steps are written as click paths in the Microsoft 365 admin consoles (no PowerShell). The 'Admin Consoles' sheet lists each console's URL and the least-privileged admin role needed. Tip: every console has a search box at the top if a menu has moved."),
    ("One exception", "Read-only Outlook on the web for unmanaged devices (item 4, step 5) has no admin-center switch; use a Defender for Cloud Apps session policy, or one Exchange PowerShell command if you do not have that licence."),
    ("Columns to fill in", "Owner and Status are for you to complete. Status values: Not started, In progress, Done, Accepted risk. Everything else is reference content."),
    ("Example", "Priority 2 > Owner: J. Smith > Status: In progress"),
    ("Criticality", "Critical = exploitable now and leads to data loss or account takeover; High = significant exposure; Medium = gaps in governance or defence in depth; Low = residual risk."),
    ("Licensing", "Most controls need Microsoft Entra ID P1 and Intune, both in Microsoft 365 Business Premium (up to 300 users) and E3+EMS / E5. Risk-based policies need Entra ID P2."),
    ("Recommended BYOD model", "Phones/tablets: Intune app protection without enrolment (MAM). Personal Windows/macOS: browser-only or Edge with Windows MAM. Full desktop apps only on Intune-compliant devices."),
    ("Caution", "Introduce Conditional Access policies in Report-only mode first and always exclude two monitored break-glass admin accounts."),
    ("Accuracy", "Microsoft renames portal menus and retires controls regularly (e.g. 'Require approved client app'). Menu paths are current as of October 2026; confirm against Microsoft Learn before changing production."),
    ("Prepared", "2026-10-06"),
]

STATUS_OPTIONS = ["Not started", "In progress", "Done", "Accepted risk"]


def rows():
    for i, item in enumerate(ITEMS, start=1):
        r = dict(item)
        r["Priority"] = i
        r["Owner"] = ""
        r["Status"] = "Not started"
        yield [r[c] for c in COLUMNS]


def build_xlsx(path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "Remediation Plan"
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    base = Font(name="Arial", size=10)

    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3864")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = border

    for values in rows():
        ws.append(values)
        r = ws.max_row
        for cell in ws[r]:
            cell.font = base
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = border
        crit = ws.cell(row=r, column=2)
        bg, fg = CRIT_FILL[crit.value]
        crit.fill = PatternFill("solid", fgColor=bg)
        crit.font = Font(name="Arial", size=10, bold=True, color=fg)
        ws.cell(row=r, column=1).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row=r, column=11).fill = PatternFill("solid", fgColor="FFFF00")
        ws.cell(row=r, column=12).fill = PatternFill("solid", fgColor="FFFF00")

    for i, w in enumerate(WIDTHS, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = ws.dimensions

    dv = DataValidation(type="list", formula1='"' + ",".join(STATUS_OPTIONS) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"L2:L{ws.max_row}")

    # Summary by criticality (formulas so it updates as Status changes)
    sm = wb.create_sheet("Summary")
    sm.append(["Criticality", "Items", "Done", "Remaining"])
    last = ws.max_row
    for crit in CRIT_FILL:
        n = sm.max_row + 1
        sm.append([
            crit,
            f"=COUNTIF('Remediation Plan'!$B$2:$B${last},A{n})",
            f"=COUNTIFS('Remediation Plan'!$B$2:$B${last},A{n},'Remediation Plan'!$L$2:$L${last},\"Done\")",
            f"=B{n}-C{n}",
        ])
    n = sm.max_row + 1
    sm.append(["Total", f"=SUM(B2:B{n-1})", f"=SUM(C2:C{n-1})", f"=SUM(D2:D{n-1})"])
    for row in sm.iter_rows():
        for cell in row:
            cell.font = Font(name="Arial", size=10, bold=cell.row in (1, n))
            cell.border = border
    for col, w in zip("ABCD", (14, 10, 10, 12)):
        sm.column_dimensions[col].width = w

    cs = wb.create_sheet("Admin Consoles")
    cs.append(["Console", "URL", "Used for", "Least-privileged role"])
    for row in CONSOLES:
        cs.append(list(row))
    for row in cs.iter_rows():
        for cell in row:
            cell.font = Font(name="Arial", size=10, bold=cell.row == 1, color="FFFFFF" if cell.row == 1 else "000000")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = border
            if cell.row == 1:
                cell.fill = PatternFill("solid", fgColor="1F3864")
    for col, w in zip("ABCD", (30, 44, 70, 50)):
        cs.column_dimensions[col].width = w

    rd = wb.create_sheet("Read Me", 0)
    for a, b in README_ROWS:
        rd.append([a, b])
    rd["A1"].font = Font(name="Arial", size=14, bold=True)
    for row in rd.iter_rows(min_row=2):
        row[0].font = Font(name="Arial", size=10, bold=True)
        row[1].font = base
        row[0].alignment = Alignment(vertical="top")
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    rd.column_dimensions["A"].width = 24
    rd.column_dimensions["B"].width = 110
    wb.active = 1
    wb.save(path)


def build_numbers(path):
    from numbers_parser import Document, RGB

    data = list(rows())
    doc = Document(
        sheet_name="Remediation Plan",
        table_name="BYOD Remediation",
        num_header_rows=1,
        num_header_cols=0,
        num_rows=len(data) + 1,
        num_cols=len(COLUMNS),
    )
    t = doc.sheets[0].tables[0]
    for c, name in enumerate(COLUMNS):
        t.write(0, c, name)
    for r, values in enumerate(data, start=1):
        for c, v in enumerate(values):
            t.write(r, c, v)

    header = doc.add_style(name="BYOD Header", font_name="Arial", font_size=11.0, bold=True,
                           font_color=RGB(255, 255, 255), bg_color=RGB(31, 56, 100), text_wrap=True)
    body = doc.add_style(name="BYOD Body", font_name="Arial", font_size=10.0, text_wrap=True, alignment=("left", "top"))
    crit_styles = {
        k: doc.add_style(name=f"BYOD {k}", font_name="Arial", font_size=10.0, bold=True, alignment=("left", "top"),
                         bg_color=RGB(*bytes.fromhex(bg)), font_color=RGB(*bytes.fromhex(fg)))
        for k, (bg, fg) in CRIT_FILL.items()
    }
    fill_in = doc.add_style(name="BYOD Fill in", font_name="Arial", font_size=10.0, alignment=("left", "top"),
                            bg_color=RGB(255, 255, 0))
    for c in range(len(COLUMNS)):
        t.set_cell_style(0, c, header)
        t.col_width(c, WIDTHS[c] * 7)
    for r, values in enumerate(data, start=1):
        for c in range(len(COLUMNS)):
            t.set_cell_style(r, c, body)
        t.set_cell_style(r, 1, crit_styles[values[1]])
        t.set_cell_style(r, 10, fill_in)
        t.set_cell_style(r, 11, fill_in)

    doc.add_sheet("Read Me", "Read Me", num_rows=len(README_ROWS), num_cols=2)
    rt = doc.sheets["Read Me"].tables[0]
    for r, (a, b) in enumerate(README_ROWS):
        rt.write(r, 0, a)
        rt.write(r, 1, b)
        rt.set_cell_style(r, 1, body)
    rt.col_width(0, 170)
    rt.col_width(1, 760)

    doc.add_sheet("Admin Consoles", "Admin Consoles", num_rows=len(CONSOLES) + 1, num_cols=4)
    ct = doc.sheets["Admin Consoles"].tables[0]
    for c, h in enumerate(["Console", "URL", "Used for", "Least-privileged role"]):
        ct.write(0, c, h)
        ct.set_cell_style(0, c, header)
    for r, row in enumerate(CONSOLES, start=1):
        for c, v in enumerate(row):
            ct.write(r, c, v)
            ct.set_cell_style(r, c, body)
    for c, w in enumerate((210, 300, 480, 340)):
        ct.col_width(c, w)
    doc.save(path)


if __name__ == "__main__":
    build_xlsx(HERE / f"{NAME}.xlsx")
    build_numbers(HERE / f"{NAME}.numbers")
    print("Wrote", NAME + ".xlsx and", NAME + ".numbers")
